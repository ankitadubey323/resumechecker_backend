import io

import pytest
from docx import Document

from app.llm import customizer
from app.llm.customizer import CustomizationOutput, LLMError, _post_check, build_candidate_text
from app.llm.schemas import CustomizationResult
from app.routers import customize
from tests.test_dashboard import DETAILS, JD, RESUME_TEXT, auth_client, make_docx  # noqa: F401

FAKE_RESULT = {
    "resume": {
        "name": "Jane Doe",
        "headline": "Backend Engineer | Python, FastAPI",
        "contact": {"email": "jane.doe@example.com", "phone": "+14155552671", "location": None, "links": []},
        "summary": "Backend developer building payment\u2011focused Python and FastAPI services.",
        "skills": [{"category": "Languages", "items": ["Python", "Kubernetes"]}],
        "experience": [
            {
                "title": "Developer",
                "company": "Acme",
                "location": None,
                "start": "2022-01",
                "end": "Present",
                "bullets": ["Built REST APIs with FastAPI", "Cut latency by 35% <b>fast</b>"],
            }
        ],
        "projects": [{"name": "Resume app", "tech": ["Next.js"], "link": None, "bullets": ["Built a resume app"]}],
        "education": [{"degree": "B.Tech", "institution": "IIT", "year": "2024", "details": None}],
        "certifications": [],
    },
    "analysis": {
        "jd_requirements": ["Python", "FastAPI", "Kubernetes"],
        "matched_keywords": ["Python", "FastAPI"],
        "missing_keywords": ["Kubernetes"],
        "bullet_changes": [
            {"section": "experience", "original": "Built APIs", "improved": "Built REST APIs with FastAPI", "reason": "Uses JD terms"}
        ],
        "suggestions": [{"area": "skills", "priority": "high", "suggestion": "Add Kubernetes if you have used it"}],
        "match_score_before": 60,
        "match_score_after": 50,
        "score_explanation": "Good Python fit; missing Kubernetes.",
    },
}


@pytest.fixture
def fake_llm(monkeypatch):
    """Replace the real Groq call; records what the endpoint sent."""
    calls: list[dict] = []

    def fake(candidate: str, job_description: str, sections: list[str] | None = None) -> CustomizationOutput:
        calls.append({"candidate": candidate, "job_description": job_description, "sections": sections})
        result = CustomizationResult.model_validate(FAKE_RESULT)
        return CustomizationOutput(result=result, model="test-model", warnings=_post_check(result, candidate))

    monkeypatch.setattr(customize, "customize_resume", fake)
    return calls


def ready_client(auth_client):  # noqa: F811
    auth_client.put("/api/profile/details", json=DETAILS)
    auth_client.post("/api/job-descriptions", json={"text": JD})
    return auth_client


def test_locked_until_profile_and_jd(auth_client, fake_llm):  # noqa: F811
    assert auth_client.get("/api/dashboard").json()["ready_for_customization"] is False
    assert auth_client.post("/api/customizations").status_code == 409

    auth_client.put("/api/profile/details", json=DETAILS)
    assert auth_client.post("/api/customizations").status_code == 409  # still no JD
    auth_client.post("/api/job-descriptions", json={"text": JD})
    assert auth_client.get("/api/dashboard").json()["ready_for_customization"] is True
    assert fake_llm == []


def test_customize_happy_path(auth_client, fake_llm):  # noqa: F811
    client = ready_client(auth_client)
    r = client.post("/api/customizations")
    assert r.status_code == 201, r.text
    body = r.json()
    assert body["model"] == "test-model" and body["source"] == "details" and body["outdated"] is False
    # The LLM got the JD and the candidate's details.
    assert fake_llm[0]["job_description"] == JD.strip()
    assert "Backend Developer" in fake_llm[0]["candidate"] and "Acme" in fake_llm[0]["candidate"]

    # Post-checks: skill marked missing is removed, invented number flagged, score not lowered.
    assert body["result"]["resume"]["skills"][0]["items"] == ["Python"]
    assert any("35" in w for w in body["warnings"])
    assert body["result"]["analysis"]["match_score_after"] == 60

    latest = client.get("/api/customizations/latest").json()
    assert latest["id"] == body["id"]
    assert client.get("/api/dashboard").json()["latest_customization"]["id"] == body["id"]


def test_resume_text_is_sent_when_uploaded(auth_client, fake_llm):  # noqa: F811
    auth_client.post("/api/resume", files={"file": ("cv.docx", make_docx(RESUME_TEXT), "x")})
    auth_client.post("/api/job-descriptions", json={"text": JD})
    body = auth_client.post("/api/customizations").json()
    assert body["source"] == "resume"
    assert "Built a resume customizer web app" in fake_llm[0]["candidate"]


def test_outdated_after_inputs_change(auth_client, fake_llm):  # noqa: F811
    client = ready_client(auth_client)
    client.post("/api/customizations")
    client.post("/api/job-descriptions", json={"text": JD + " Also Kubernetes."})
    assert client.get("/api/customizations/latest").json()["outdated"] is True
    assert client.get("/api/dashboard").json()["latest_customization"]["outdated"] is True


def test_llm_errors_are_user_friendly(auth_client, monkeypatch):  # noqa: F811
    client = ready_client(auth_client)

    def boom(candidate, jd, sections=None):
        raise LLMError("The AI service is busy right now (free-tier limit). Please try again in a minute.", 429)

    monkeypatch.setattr(customize, "customize_resume", boom)
    r = client.post("/api/customizations")
    assert r.status_code == 429 and "busy" in r.json()["detail"]
    assert client.get("/api/customizations/latest").status_code == 404  # nothing saved


def test_missing_api_key(auth_client, monkeypatch):  # noqa: F811
    client = ready_client(auth_client)
    monkeypatch.setattr(customizer.get_settings(), "groq_api_key", "")
    r = client.post("/api/customizations")
    assert r.status_code == 503 and "GROQ_API_KEY" in r.json()["detail"]


def test_per_user_hourly_limit(auth_client, fake_llm, monkeypatch):  # noqa: F811
    client = ready_client(auth_client)
    monkeypatch.setattr(customize.get_settings(), "customizations_per_hour", 2)
    assert client.post("/api/customizations").status_code == 201
    assert client.post("/api/customizations").status_code == 201
    assert client.post("/api/customizations").status_code == 429


@pytest.mark.parametrize("fmt,ctype,magic", [("pdf", "application/pdf", b"%PDF-"), ("docx", "officedocument", b"PK"), ("txt", "text/plain", b"JANE DOE")])
def test_downloads(auth_client, fake_llm, fmt, ctype, magic):  # noqa: F811
    client = ready_client(auth_client)
    cid = client.post("/api/customizations").json()["id"]
    r = client.get(f"/api/customizations/{cid}/download?format={fmt}")
    assert r.status_code == 200
    assert ctype in r.headers["content-type"]
    assert r.headers["content-disposition"] == f'attachment; filename="Jane_Doe_Resume.{fmt}"'
    assert r.content.startswith(magic)
    if fmt == "docx":
        text = "\n".join(p.text for p in Document(io.BytesIO(r.content)).paragraphs)
        assert "Built REST APIs with FastAPI" in text and "<b>fast</b>" in text  # escaped, kept literally


def test_download_bad_format_and_other_users(client, auth_client, fake_llm):  # noqa: F811
    c = ready_client(auth_client)
    cid = c.post("/api/customizations").json()["id"]
    assert c.get(f"/api/customizations/{cid}/download?format=exe").status_code == 422
    assert c.get("/api/customizations/not-an-id/download").status_code == 404

    c.post("/api/auth/logout")
    c.cookies.clear()
    c.post(
        "/api/auth/signup",
        json={"full_name": "Eve Other", "email": "eve@example.com", "phone": "+14155552672", "password": "Str0ng!Pass", "confirm_password": "Str0ng!Pass"},
    )
    assert c.get(f"/api/customizations/{cid}/download").status_code == 404


def test_unusual_characters_are_normalized():
    result = customizer._normalize_text(CustomizationResult.model_validate(FAKE_RESULT))
    assert result.resume.summary == "Backend developer building payment-focused Python and FastAPI services."


def test_selected_input_decides_what_is_sent():
    from app.services.dashboard import input_source

    user = {"full_name": "Jane", "email": "j@example.com", "phone": "+1"}
    profile = {"details_updated_at": 1, **DETAILS, "selected_input": "resume"}
    resume = {"text": "RESUME BODY", "uploaded_at": 1}

    assert input_source(profile, resume) == "resume"
    text = build_candidate_text(user, profile, resume)
    assert "RESUME BODY" in text and "Acme" not in text

    profile["selected_input"] = "details"
    assert input_source(profile, resume) == "details"
    text = build_candidate_text(user, profile, resume)
    assert "RESUME BODY" not in text and "Work experience: Developer at Acme (2022-01 to Present)" in text

    # Selected input missing -> fall back to the one that exists.
    assert input_source(profile, None) == "details"
    assert input_source({"selected_input": "details"}, resume) == "resume"
    assert input_source(None, None) is None


def test_each_input_has_its_own_result(auth_client, fake_llm):  # noqa: F811
    c = auth_client
    c.post("/api/resume", files={"file": ("cv.docx", make_docx(RESUME_TEXT), "x")})
    c.post("/api/job-descriptions", json={"text": JD})
    assert c.get("/api/dashboard").json()["customization_input"] == "resume"
    from_resume = c.post("/api/customizations").json()
    assert from_resume["source"] == "resume"
    assert "Built a resume customizer web app" in fake_llm[-1]["candidate"]

    # Saving details selects them; the resume result is kept but not shown for this input.
    d = c.put("/api/profile/details", json=DETAILS).json()
    assert d["customization_input"] == "details" and d["latest_customization"] is None
    from_details = c.post("/api/customizations").json()
    assert from_details["source"] == "details"
    assert "Backend Developer" in fake_llm[-1]["candidate"] and "Built a resume customizer" not in fake_llm[-1]["candidate"]

    # Switching back (choosing the tab) shows the resume result again, not regenerated.
    d = c.put("/api/profile/input", json={"input": "resume"}).json()
    assert d["customization_input"] == "resume" and d["latest_customization"]["id"] == from_resume["id"]
    assert c.get("/api/customizations/latest").json()["id"] == from_resume["id"]
    assert c.get("/api/customizations/latest?source=details").json()["id"] == from_details["id"]
    assert len(fake_llm) == 2

    # Saving details again only makes the details result outdated.
    c.put("/api/profile/details", json={**DETAILS, "headline": "Senior Backend Developer"})
    assert c.get("/api/customizations/latest?source=details").json()["outdated"] is True
    assert c.get("/api/customizations/latest?source=resume").json()["outdated"] is False


def test_select_input_requires_that_input(auth_client):  # noqa: F811
    assert auth_client.put("/api/profile/input", json={"input": "resume"}).status_code == 409
    assert auth_client.put("/api/profile/input", json={"input": "details"}).status_code == 409
    assert auth_client.put("/api/profile/input", json={"input": "other"}).status_code == 422


def test_invented_experience_is_removed():
    data = {**FAKE_RESULT, "resume": {**FAKE_RESULT["resume"], "experience": [
        {"title": "Developer", "company": "Acme", "location": None, "start": "2022-01", "end": "Present", "bullets": ["Built APIs"]},
        {"title": "Java Developer (Academic Project)", "company": "SRM University", "location": None, "start": "2024", "end": "2025", "bullets": ["Built a library app"]},
    ]}}
    result = CustomizationResult.model_validate(data)
    candidate = "Work experience: Developer at Acme (2022-01 to Present)\nEducation: B.Tech, SRM University, 2025"
    warnings = _post_check(result, candidate)
    assert [e.company for e in result.resume.experience] == ["Acme"]  # 2024 isn't in the data
    assert any("SRM University" in w for w in warnings)


# --- bonus features: versions, sections, ATS, diff texts, cover letter -------------


def test_sections_are_passed_and_saved(auth_client, fake_llm):  # noqa: F811
    c = ready_client(auth_client)
    body = c.post("/api/customizations", json={"sections": ["skills", "experience"]}).json()
    assert body["sections"] == ["skills", "experience"]
    assert fake_llm[-1]["sections"] == ["skills", "experience"]
    # No body -> all sections.
    assert c.post("/api/customizations").json()["sections"] == ["summary", "skills", "experience", "projects"]
    assert c.post("/api/customizations", json={"sections": []}).status_code == 422
    assert c.post("/api/customizations", json={"sections": ["education"]}).status_code == 422


def test_scope_instruction_lists_kept_sections():
    from app.llm.prompts import sections_instruction

    assert sections_instruction(["summary", "skills", "experience", "projects"]) == "SECTION SCOPE: tailor all sections."
    text = sections_instruction(["skills"])
    assert "TAILOR these sections fully, following all the writing rules: skills." in text
    assert "DO NOT TAILOR these sections: summary, experience, projects." in text
    assert "no summary" in text  # guidance for a kept-but-missing summary


def test_versions_list_and_get_by_id(auth_client, fake_llm):  # noqa: F811
    c = ready_client(auth_client)
    first = c.post("/api/customizations").json()
    second = c.post("/api/customizations", json={"sections": ["summary"]}).json()
    versions = c.get("/api/customizations?source=details").json()
    assert [v["id"] for v in versions] == [second["id"], first["id"]]
    assert versions[0]["sections"] == ["summary"] and isinstance(versions[0]["ats_score"], int)
    assert c.get("/api/customizations?source=resume").json() == []

    got = c.get(f"/api/customizations/{first['id']}").json()
    assert got["id"] == first["id"]
    # Diff view inputs: what was sent to the AI, and the tailored resume as text.
    assert "Backend Developer" in got["input_text"] and got["tailored_text"].startswith("JANE DOE")
    assert c.get("/api/customizations/000000000000000000000000").status_code == 404


def test_ats_report_scores_and_explains():
    from app.services.ats import ats_report

    result = CustomizationResult.model_validate(FAKE_RESULT)
    report = ats_report(result.resume, result.analysis)
    assert 0 <= report.score <= 100
    labels = {c.label: c for c in report.checks}
    assert set(labels) == {"Job keywords", "Essential sections", "Action verbs", "Bullet length", "Quantified results", "No placeholders"}
    assert sum(c.max_points for c in report.checks) == 100

    result.resume.experience[0].bullets[0] = "Reduced costs by [X]%"
    report = ats_report(result.resume, result.analysis)
    assert not {c.label: c for c in report.checks}["No placeholders"].passed


def test_cover_letter_streams_saves_and_downloads(auth_client, fake_llm, monkeypatch):  # noqa: F811
    c = ready_client(auth_client)
    cid = c.post("/api/customizations").json()["id"]
    seen = {}

    def fake_stream(resume_text, jd, missing):
        seen.update(resume=resume_text, missing=missing)
        return iter(["Dear Hiring Manager,\n\n", "I build APIs.", "\n\nSincerely,\nJane Doe"])

    monkeypatch.setattr(customize, "stream_cover_letter", fake_stream)
    r = c.post(f"/api/customizations/{cid}/cover-letter")
    assert r.status_code == 200 and r.headers["content-type"].startswith("text/plain")
    assert "no-transform" in r.headers["cache-control"]  # otherwise proxies gzip-buffer the stream
    assert r.text == "Dear Hiring Manager,\n\nI build APIs.\n\nSincerely,\nJane Doe"
    assert "JANE DOE" in seen["resume"] and seen["missing"] == ["Kubernetes"]
    assert c.get(f"/api/customizations/{cid}").json()["cover_letter"]["text"].endswith("Jane Doe")

    for fmt, magic in (("pdf", b"%PDF-"), ("docx", b"PK"), ("txt", b"Jane Doe")):
        d = c.get(f"/api/customizations/{cid}/cover-letter/download?format={fmt}")
        assert d.status_code == 200 and d.content.startswith(magic)
        assert d.headers["content-disposition"] == f'attachment; filename="Jane_Doe_Cover_Letter.{fmt}"'


def test_cover_letter_errors(auth_client, fake_llm, monkeypatch):  # noqa: F811
    c = ready_client(auth_client)
    cid = c.post("/api/customizations").json()["id"]
    assert c.get(f"/api/customizations/{cid}/cover-letter/download").status_code == 404  # not generated yet

    def busy(*a):
        raise LLMError("The AI service is busy right now (free-tier limit). Please try again in a minute.", 429)

    monkeypatch.setattr(customize, "stream_cover_letter", busy)
    r = c.post(f"/api/customizations/{cid}/cover-letter")
    assert r.status_code == 429 and "busy" in r.json()["detail"]

    def breaks_midway(*a):
        def gen():
            yield "Dear Hiring Manager,"
            raise RuntimeError("connection dropped")

        return gen()

    monkeypatch.setattr(customize, "stream_cover_letter", breaks_midway)
    r = c.post(f"/api/customizations/{cid}/cover-letter")
    assert customize.STREAM_ERROR_MARKER in r.text
    assert c.get(f"/api/customizations/{cid}").json()["cover_letter"] is None  # partial letter not saved
