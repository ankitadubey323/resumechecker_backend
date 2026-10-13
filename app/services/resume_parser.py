"""Extract plain text from an uploaded PDF or DOCX resume."""

import io
import re
import zipfile
from dataclasses import dataclass
from pathlib import PurePath

import pdfplumber
from docx import Document

MAX_RESUME_BYTES = 5 * 1024 * 1024
MAX_PAGES = 10
MAX_DOCX_UNCOMPRESSED = 30 * 1024 * 1024  # guards against zip bombs
MAX_TEXT_CHARS = 50_000
MIN_TEXT_CHARS = 50

PDF_MIME = "application/pdf"
DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


class ResumeParseError(ValueError):
    def __init__(self, message: str, status_code: int = 422):
        super().__init__(message)
        self.status_code = status_code


@dataclass
class ParsedResume:
    text: str
    content_type: str
    page_count: int | None


def parse_resume(filename: str, data: bytes) -> ParsedResume:
    if not data:
        raise ResumeParseError("The file is empty.")
    if len(data) > MAX_RESUME_BYTES:
        raise ResumeParseError("The file is larger than 5 MB.", 413)

    # Trust the file's bytes, not the name or the browser-supplied MIME type.
    ext = PurePath(filename or "").suffix.lower()
    if ext == ".pdf" and data.startswith(b"%PDF-"):
        text, pages, mime = *_parse_pdf(data), PDF_MIME
    elif ext == ".docx" and data.startswith(b"PK\x03\x04"):
        text, pages, mime = _parse_docx(data), None, DOCX_MIME
    else:
        raise ResumeParseError("Only PDF or DOCX files are supported.", 415)

    text = _clean(text)
    if len(text) < MIN_TEXT_CHARS:
        raise ResumeParseError(
            "We couldn't find readable text in this file. If it's a scanned PDF, "
            "upload a text-based PDF or a DOCX instead."
        )
    return ParsedResume(text=text[:MAX_TEXT_CHARS], content_type=mime, page_count=pages)


def _parse_pdf(data: bytes) -> tuple[str, int]:
    try:
        with pdfplumber.open(io.BytesIO(data)) as pdf:
            if len(pdf.pages) > MAX_PAGES:
                raise ResumeParseError(f"Resumes can be at most {MAX_PAGES} pages.")
            return "\n".join(page.extract_text() or "" for page in pdf.pages), len(pdf.pages)
    except ResumeParseError:
        raise
    except Exception:
        raise ResumeParseError("We couldn't read this PDF. It may be corrupted or password-protected.")


def _parse_docx(data: bytes) -> str:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as zf:
            if "word/document.xml" not in zf.namelist():
                raise ResumeParseError("This doesn't look like a valid Word document.")
            if sum(i.file_size for i in zf.infolist()) > MAX_DOCX_UNCOMPRESSED:
                raise ResumeParseError("This Word document is too large to process.")
        doc = Document(io.BytesIO(data))
    except ResumeParseError:
        raise
    except Exception:
        raise ResumeParseError("We couldn't read this Word document. It may be corrupted.")

    lines = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            lines.append(" | ".join(dict.fromkeys(c.text.strip() for c in row.cells if c.text.strip())))
    return "\n".join(lines)


def _clean(text: str) -> str:
    text = text.replace("\x00", "")
    lines = [re.sub(r"[ \t]+", " ", line).strip() for line in text.splitlines()]
    return re.sub(r"\n{3,}", "\n\n", "\n".join(lines)).strip()
