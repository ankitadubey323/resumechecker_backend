"""Render a tailored resume (app.llm.schemas.TailoredResume) as TXT, DOCX or PDF.

All three use the same simple, single-column, ATS-friendly layout:
name, headline, contact line, then Summary / Skills / Experience / Projects /
Education / Certifications.
"""

import io
import re
from xml.sax.saxutils import escape

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import HRFlowable, ListFlowable, ListItem, Paragraph, SimpleDocTemplate, Spacer

from app.llm.schemas import TailoredResume

MEDIA_TYPES = {
    "pdf": "application/pdf",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "txt": "text/plain; charset=utf-8",
}


def filename_for(resume: TailoredResume, fmt: str) -> str:
    safe = re.sub(r"[^A-Za-z0-9]+", "_", resume.name).strip("_") or "Resume"
    return f"{safe}_Resume.{fmt}"


def _contact_line(r: TailoredResume) -> str:
    parts = [r.contact.email, r.contact.phone, r.contact.location, *r.contact.links]
    return "  |  ".join(p for p in parts if p)


def _dates(start: str, end: str) -> str:
    return f"{start} – {end}" if start or end else ""


# --- plain text --------------------------------------------------------------


def to_txt(r: TailoredResume) -> bytes:
    out = [r.name.upper(), r.headline, _contact_line(r), ""]

    def section(title: str) -> None:
        out.extend([title.upper(), "-" * len(title)])

    section("Summary")
    out += [r.summary, ""]
    if r.skills:
        section("Skills")
        out += [f"{g.category}: {', '.join(g.items)}" for g in r.skills] + [""]
    if r.experience:
        section("Experience")
        for e in r.experience:
            loc = f", {e.location}" if e.location else ""
            out.append(f"{e.title} — {e.company}{loc}  ({_dates(e.start, e.end)})")
            out += [f"  • {b}" for b in e.bullets] + [""]
    if r.projects:
        section("Projects")
        for p in r.projects:
            tech = f" ({', '.join(p.tech)})" if p.tech else ""
            out.append(f"{p.name}{tech}" + (f" — {p.link}" if p.link else ""))
            out += [f"  • {b}" for b in p.bullets] + [""]
    if r.education:
        section("Education")
        for ed in r.education:
            extra = "".join(f", {x}" for x in (ed.year, ed.details) if x)
            out.append(f"{ed.degree}, {ed.institution}{extra}")
        out.append("")
    if r.certifications:
        section("Certifications")
        out += [f"  • {c}" for c in r.certifications]
    return "\n".join(out).strip().encode() + b"\n"


# --- DOCX --------------------------------------------------------------------


def to_docx(r: TailoredResume) -> bytes:
    doc = Document()
    for s in doc.sections:
        s.top_margin = s.bottom_margin = Pt(42)
        s.left_margin = s.right_margin = Pt(50)
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)

    def para(text: str = "", *, size: float | None = None, bold=False, center=False, color=None, space_after=2):
        p = doc.add_paragraph()
        run = p.add_run(text)
        run.bold = bold
        if size:
            run.font.size = Pt(size)
        if color:
            run.font.color.rgb = RGBColor(*color)
        if center:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(space_after)
        return p

    def heading(title: str) -> None:
        p = para(title.upper(), size=11, bold=True, color=(0x31, 0x2E, 0x81), space_after=3)
        p.paragraph_format.space_before = Pt(8)

    def bullets(items: list[str]) -> None:
        for b in items:
            p = doc.add_paragraph(b, style="List Bullet")
            p.paragraph_format.space_after = Pt(1)

    para(r.name, size=20, bold=True, center=True, space_after=0)
    para(r.headline, size=11, center=True, color=(0x52, 0x52, 0x5B), space_after=0)
    para(_contact_line(r), size=9, center=True, color=(0x52, 0x52, 0x5B), space_after=6)

    heading("Summary")
    para(r.summary)
    if r.skills:
        heading("Skills")
        for g in r.skills:
            p = doc.add_paragraph()
            p.add_run(f"{g.category}: ").bold = True
            p.add_run(", ".join(g.items))
            p.paragraph_format.space_after = Pt(1)
    if r.experience:
        heading("Experience")
        for e in r.experience:
            p = doc.add_paragraph()
            p.add_run(f"{e.title} — {e.company}").bold = True
            if e.location:
                p.add_run(f", {e.location}")
            p.add_run(f"   {_dates(e.start, e.end)}").italic = True
            p.paragraph_format.space_after = Pt(1)
            bullets(e.bullets)
    if r.projects:
        heading("Projects")
        for pr in r.projects:
            p = doc.add_paragraph()
            p.add_run(pr.name).bold = True
            if pr.tech:
                p.add_run(f" ({', '.join(pr.tech)})")
            if pr.link:
                p.add_run(f" — {pr.link}")
            p.paragraph_format.space_after = Pt(1)
            bullets(pr.bullets)
    if r.education:
        heading("Education")
        for ed in r.education:
            extra = "".join(f", {x}" for x in (ed.year, ed.details) if x)
            para(f"{ed.degree}, {ed.institution}{extra}", space_after=1)
    if r.certifications:
        heading("Certifications")
        bullets(r.certifications)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# --- PDF ---------------------------------------------------------------------


def to_pdf(r: TailoredResume) -> bytes:
    styles = getSampleStyleSheet()
    accent = colors.HexColor("#312e81")
    muted = colors.HexColor("#52525b")
    st = {
        "name": ParagraphStyle("name", parent=styles["Title"], fontSize=20, leading=24, spaceAfter=2),
        "headline": ParagraphStyle("headline", parent=styles["Normal"], alignment=TA_CENTER, textColor=muted, fontSize=10.5),
        "contact": ParagraphStyle("contact", parent=styles["Normal"], alignment=TA_CENTER, textColor=muted, fontSize=8.5, spaceAfter=6),
        "h": ParagraphStyle("h", parent=styles["Heading2"], fontSize=11, textColor=accent, spaceBefore=8, spaceAfter=2),
        "body": ParagraphStyle("body", parent=styles["Normal"], fontSize=9.5, leading=12.5),
        "item": ParagraphStyle("item", parent=styles["Normal"], fontSize=9.5, leading=12.5, spaceBefore=4),
    }

    def p(text: str, style: str) -> Paragraph:
        return Paragraph(escape(text), st[style])

    def heading(title: str) -> list:
        return [Paragraph(title.upper(), st["h"]), HRFlowable(width="100%", thickness=0.6, color=accent, spaceAfter=3)]

    def bullet_list(items: list[str]) -> ListFlowable:
        return ListFlowable(
            [ListItem(p(b, "body"), leftIndent=10) for b in items], bulletType="bullet", start="•", leftIndent=10, bulletFontSize=8
        )

    story: list = [p(r.name, "name"), p(r.headline, "headline"), p(_contact_line(r), "contact")]
    story += heading("Summary") + [p(r.summary, "body")]
    if r.skills:
        story += heading("Skills")
        story += [Paragraph(f"<b>{escape(g.category)}:</b> {escape(', '.join(g.items))}", st["body"]) for g in r.skills]
    if r.experience:
        story += heading("Experience")
        for e in r.experience:
            loc = f", {escape(e.location)}" if e.location else ""
            story.append(
                Paragraph(
                    f"<b>{escape(e.title)} — {escape(e.company)}</b>{loc}"
                    f"  <font color='#52525b'><i>{escape(_dates(e.start, e.end))}</i></font>",
                    st["item"],
                )
            )
            story.append(bullet_list(e.bullets))
    if r.projects:
        story += heading("Projects")
        for pr in r.projects:
            tech = f" ({escape(', '.join(pr.tech))})" if pr.tech else ""
            link = f" — {escape(pr.link)}" if pr.link else ""
            story.append(Paragraph(f"<b>{escape(pr.name)}</b>{tech}{link}", st["item"]))
            story.append(bullet_list(pr.bullets))
    if r.education:
        story += heading("Education")
        for ed in r.education:
            extra = "".join(f", {escape(x)}" for x in (ed.year, ed.details) if x)
            story.append(Paragraph(f"<b>{escape(ed.degree)}</b>, {escape(ed.institution)}{extra}", st["body"]))
    if r.certifications:
        story += heading("Certifications") + [bullet_list(r.certifications)]

    buf = io.BytesIO()
    SimpleDocTemplate(
        buf, pagesize=A4, leftMargin=16 * mm, rightMargin=16 * mm, topMargin=14 * mm, bottomMargin=14 * mm,
        title=f"{r.name} - Resume", author=r.name,
    ).build(story)
    return buf.getvalue()


EXPORTERS = {"pdf": to_pdf, "docx": to_docx, "txt": to_txt}


# --- cover letter --------------------------------------------------------------


def cover_letter_filename(resume: TailoredResume, fmt: str) -> str:
    return filename_for(resume, fmt).replace("_Resume.", "_Cover_Letter.")


def _paragraphs(text: str) -> list[str]:
    return [p.strip() for p in re.split(r"\n\s*\n", text.strip()) if p.strip()]


def cover_letter_txt(r: TailoredResume, text: str) -> bytes:
    return f"{r.name}\n{_contact_line(r)}\n\n{text.strip()}\n".encode()


def cover_letter_docx(r: TailoredResume, text: str) -> bytes:
    doc = Document()
    for s in doc.sections:
        s.top_margin = s.bottom_margin = Pt(56)
        s.left_margin = s.right_margin = Pt(64)
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)
    head = doc.add_paragraph()
    head.add_run(r.name).bold = True
    head.runs[0].font.size = Pt(16)
    contact = doc.add_paragraph(_contact_line(r))
    contact.runs[0].font.size = Pt(9)
    contact.paragraph_format.space_after = Pt(18)
    for para in _paragraphs(text):
        p = doc.add_paragraph()
        lines = para.split("\n")
        for i, line in enumerate(lines):
            p.add_run(line)
            if i < len(lines) - 1:
                p.add_run().add_break()
        p.paragraph_format.space_after = Pt(10)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def cover_letter_pdf(r: TailoredResume, text: str) -> bytes:
    styles = getSampleStyleSheet()
    body = ParagraphStyle("body", parent=styles["Normal"], fontSize=10.5, leading=15, spaceAfter=9)
    story: list = [
        Paragraph(f"<b>{escape(r.name)}</b>", ParagraphStyle("n", parent=styles["Normal"], fontSize=16, leading=20)),
        Paragraph(escape(_contact_line(r)), ParagraphStyle("c", parent=styles["Normal"], fontSize=8.5, textColor=colors.HexColor("#52525b"))),
        Spacer(1, 16),
    ]
    story += [Paragraph(escape(p).replace("\n", "<br/>"), body) for p in _paragraphs(text)]
    buf = io.BytesIO()
    SimpleDocTemplate(
        buf, pagesize=A4, leftMargin=22 * mm, rightMargin=22 * mm, topMargin=20 * mm, bottomMargin=20 * mm,
        title=f"{r.name} - Cover Letter", author=r.name,
    ).build(story)
    return buf.getvalue()


COVER_LETTER_EXPORTERS = {"pdf": cover_letter_pdf, "docx": cover_letter_docx, "txt": cover_letter_txt}
